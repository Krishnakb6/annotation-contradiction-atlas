"""
Enhanced script to generate a publication-quality IEEE-style Word document (.docx)
from docs/research_paper.md using python-docx.
Includes proper IEEE citation resolution ([1], [2]), equation rendering,
custom styled tables with padding, and callout abstract boxes.
"""

import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

CITE_MAP = {
    'uniprot2023': 1,
    'uniprot2025': 2,
    'go2023': 3,
    'brenner1999': 4,
    'devos2001': 5,
    'schnoes2009': 6,
    'jones2007': 7,
    'manning2002': 8,
    'eyers2013': 9,
    'murphy2014': 10,
    'eyers2016': 11,
    'jacobsen2019': 12,
    'hanks2001': 13,
    'kornev2008': 14,
    'taylor2012': 15,
    'scheeff2009': 16,
    'fukuda2003': 17,
    'wickstrom2010': 18,
    'lin2023': 19,
    'rives2021': 20,
    'elnaggar2021': 21,
    'brandes2022': 22,
    'needleman1970': 23,
    'altschul1990': 24,
    'cock2009': 25,
}

def set_cell_background(cell, hex_color):
    """Set the background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set inner padding for table cell in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3", sz="4", val="single"):
    """Set clean subtle borders on table."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def replace_citations(text):
    """Convert \\cite{key1, key2} to [1], [2]."""
    def _cite_repl(m):
        keys = [k.strip() for k in m.group(1).split(',')]
        nums = [f"[{CITE_MAP.get(k, k)}]" for k in keys]
        return ", ".join(nums)
    return re.sub(r'\\cite\{([^}]+)\}', _cite_repl, text)

def clean_latex_math(text):
    """Convert LaTeX math syntax to clean readable Unicode academic notation."""
    # First resolve citations
    text = replace_citations(text)
    
    # Standard math macro replacements
    math_map = [
        (r'\mathbb{R}^{320}', 'ℝ³²⁰'),
        (r'\mathbb{R}', 'ℝ'),
        (r'\tau_{\text{sim}}', 'τ_sim'),
        (r'\tau_{\text{emb}}', 'τ_emb'),
        (r'\tau_{\text{feat}}', 'τ_feat'),
        (r'\tau_{\text{priority}}', 'τ_priority'),
        (r'\tau', 'τ'),
        (r'\Delta_{\text{class}}', 'Δ_class'),
        (r'\Delta', 'Δ'),
        (r'\alpha', 'α'),
        (r'\beta', 'β'),
        (r'\gamma', 'γ'),
        (r'\approx', '≈'),
        (r'\le', '≤'),
        (r'\ge', '≥'),
        (r'\times', '×'),
        (r'\cdot', '·'),
        (r'\sum', '∑'),
        (r'\in', '∈'),
        (r'\to', '→'),
        (r'\dots', '…'),
        (r'\min', 'min'),
        (r'\max', 'max'),
        (r'\cap', '∩'),
        (r'\cup', '∪'),
        (r'\mathbf{e}_Q', 'e_Q'),
        (r'\mathbf{e}_i', 'e_i'),
        (r'\mathbf{e}_j', 'e_j'),
        (r'\mathbf{e}', 'e'),
        (r'\mathbf{c}_C', 'c_C'),
        (r'\mathbf{c}', 'c'),
        (r'\mathbf{v}_w', 'v_w'),
        (r'\mathbf{v}', 'v'),
        (r'\mathbf{E}_C', 'E_C'),
        (r'\mathbf{E}_{\text{own}}', 'E_own'),
        (r'\mathbf{E}_{\text{other}}', 'E_other'),
        (r'\mathbf{E}', 'E'),
        (r'\mathcal{N}_k', 'N_k'),
        (r'\mathcal{N}', 'N'),
        (r'\|\mathbf{v}\|_2', '||v||₂'),
        (r'\|\sum_{i=1}^N \mathbf{e}_i\|_2', '||∑ e_i||₂'),
        (r'\|\sum_{i=1}^{N_C} \mathbf{e}_i\|_2', '||∑ e_i||₂'),
        (r'\|\mathbf{u}\|_2', '||u||₂'),
        (r'\|\mathbf{v}\|_2', '||v||₂'),
        (r'\text{align}', 'align'),
        (r'\text{kmer}', 'kmer'),
        (r'\text{ident}', 'ident'),
        (r'\text{score}', 'score'),
        (r'\text{base}', 'base'),
        (r'\text{feat}', 'feat'),
        (r'\text{prov}', 'prov'),
        (r'\text{emb}', 'emb'),
        (r'\text{sim}', 'sim'),
        (r'\text{pos}', 'pos'),
        (r'\text{canon}', 'canon'),
        (r'\text{variant}', 'variant'),
        (r'\text{centroid}', 'centroid'),
        (r'\text{knn}', 'knn'),
        (r'\text{combined}', 'combined'),
        (r'\text{cos}', 'cos'),
        (r'\text{own}', 'own'),
        (r'\text{other}', 'other'),
        (r'\text{cat}', 'cat'),
        (r'\text{Precision}', 'Precision'),
        (r'\text{Recall}', 'Recall'),
        (r'\text{Specificity}', 'Specificity'),
        (r'\text{Accuracy}', 'Accuracy'),
        (r'\text{Condition: }', 'Condition: '),
        (r'\text{Span Constraints: }', 'Span Constraints: '),
        (r'\text{Alignment Parameters: }', 'Alignment Parameters: '),
        (r'\text{Match}', 'Match'),
        (r'\text{Mismatch}', 'Mismatch'),
        (r'\text{Gap Open}', 'Gap Open'),
        (r'\text{Gap Extend}', 'Gap Extend'),
        (r'\text{Asp}^{166}', 'Asp166'),
        (r'\text{Ala}', 'Ala'),
        (r'\text{Lys}^{72}', 'Lys72'),
        (r'\text{Glu}^{91}', 'Glu91'),
        (r'\text{Asp}^{184}', 'Asp184'),
        (r'\text{Mg}^{2+}', 'Mg²⁺'),
        (r'\text{V-A-I-K}', 'V-A-I-K'),
        (r'\text{A-X-K}', 'A-X-K'),
        (r'\text{H-R-D-x-[KR]-x-x-N}', 'H-R-D-x-[KR]-x-x-N'),
        (r'\text{H-R-D}', 'H-R-D'),
        (r'\text{D-F-G}', 'D-F-G'),
        (r'\text{H-G-A}', 'H-G-A'),
        (r'\text{S-S-G}', 'S-S-G'),
        (r'\text{F-K-K}', 'F-K-K'),
        (r'\text{D-F-F}', 'D-F-F'),
        (r'\text{HIGH}', 'HIGH'),
        (r'\text{MODERATE}', 'MODERATE'),
        (r'\text{LOW}', 'LOW'),
        (r'\text{Priority}', 'Priority'),
        (r'\text{Score}', 'Score'),
        (r'\text{Identity}', 'Identity'),
        (r'\text{Coverage}', 'Coverage'),
        (r'\text{Similarity}', 'Similarity'),
        (r'\text{Penalty}', 'Penalty'),
        (r'\text{Positives}', 'Positives'),
        (r'\text{Negatives}', 'Negatives'),
        (r'\text{Swaps}', 'Swaps'),
        (r'\text{Controls}', 'Controls'),
        (r'\text{Knockouts}', 'Knockouts'),
        (r'\left(', '('),
        (r'\right)', ')'),
        (r'\left\{', '{'),
        (r'\right\}', '}'),
        (r'\left[', '['),
        (r'\right]', ']'),
        (r'\quad', '   '),
        (r'\qquad', '      '),
        (r'\{', '{'),
        (r'\}', '}'),
        (r'\_', '_'),
    ]
    for pattern, repl in math_map:
        text = text.replace(pattern, repl)
    
    # Clean fractions: \frac{a}{b} -> (a / b)
    text = re.sub(r'\\frac\{([^}]+)\}\{([^}]+)\}', r'(\1 / \2)', text)
    # Clean square roots: \sqrt{x} -> √(x)
    text = re.sub(r'\\sqrt\{([^}]+)\}', r'√(\1)', text)
    # Clean remaining \text{...}
    text = re.sub(r'\\text\{([^}]+)\}', r'\1', text)
    # Remove cases environment markers
    text = text.replace(r'\begin{cases}', '').replace(r'\end{cases}', '')
    # Strip math delimiters
    text = text.replace('$$', '')
    text = text.replace('$', '')
    return text

def add_formatted_runs(paragraph, text, italic_all=False, bold_all=False):
    """Parse inline markdown (bold, italic, code) and add styled runs."""
    text = clean_latex_math(text)
    
    tokens = re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**'):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            if italic_all:
                run.italic = True
        elif token.startswith('*') and token.endswith('*'):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
            if bold_all:
                run.bold = True
        elif token.startswith('`') and token.endswith('`'):
            run = paragraph.add_run(token[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(160, 40, 40)
        else:
            run = paragraph.add_run(token)
            if italic_all:
                run.italic = True
            if bold_all:
                run.bold = True

def parse_markdown_table(lines):
    """Parse markdown table lines into headers and row data."""
    headers = []
    rows = []
    for line in lines:
        line = line.strip()
        if not line or not line.startswith('|'):
            continue
        parts = [p.strip() for p in line.split('|')[1:-1]]
        if not headers:
            headers = parts
        elif all(set(p).issubset({'-', ':', ' '}) for p in parts):
            continue  # Separator line
        else:
            rows.append(parts)
    return headers, rows

def create_document():
    doc = Document()
    
    # 1. Page setup - Standard Margins (1 inch)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.different_first_page_header_footer = True
        
        # Running Header (pages 2+)
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("CHARAN et al.: THE ANNOTATION CONTRADICTION ATLAS FOR HUMAN PROTEIN KINASES")
        hrun.font.name = 'Times New Roman'
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(120, 120, 120)
        
        # Footer
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("Page ")
        frun.font.name = 'Times New Roman'
        frun.font.size = Pt(9)
        fldSimple = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
        fp._p.append(fldSimple)
        frun2 = fp.add_run(" of ")
        frun2.font.name = 'Times New Roman'
        frun2.font.size = Pt(9)
        fldNumPages = parse_xml(r'<w:fldSimple %s w:instr="NUMPAGES"/>' % nsdecls('w'))
        fp._p.append(fldNumPages)

    # 2. Base Normal Style
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(30, 30, 30)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(4)

    # 3. Read markdown source
    md_path = os.path.join('docs', 'research_paper.md')
    with open(md_path, 'r', encoding='utf-8') as f:
        md_text = f.read()

    lines = md_text.split('\n')
    
    i = 0
    in_code_block = False
    code_lines = []
    in_table = False
    table_lines = []
    
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        
        # Handle code blocks
        if stripped.startswith('```'):
            if in_code_block:
                in_code_block = False
                tbl = doc.add_table(rows=1, cols=1)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell = tbl.cell(0, 0)
                set_cell_background(cell, "F4F5F7")
                set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
                set_table_borders(tbl, color="D0D5DD", sz="6", val="single")
                cp = cell.paragraphs[0]
                cp.paragraph_format.line_spacing = 1.0
                cp.paragraph_format.space_after = Pt(0)
                code_text = '\n'.join(code_lines)
                c_run = cp.add_run(code_text)
                c_run.font.name = 'Consolas'
                c_run.font.size = Pt(8.5)
                c_run.font.color.rgb = RGBColor(33, 37, 41)
                code_lines = []
                sp = doc.add_paragraph()
                sp.paragraph_format.space_after = Pt(4)
                sp.paragraph_format.line_spacing = 1.0
            else:
                in_code_block = True
                code_lines = []
            i += 1
            continue
            
        if in_code_block:
            code_lines.append(line)
            i += 1
            continue
            
        # Handle markdown tables
        if stripped.startswith('|') and '|' in stripped[1:]:
            if not in_table:
                in_table = True
                table_lines = [line]
            else:
                table_lines.append(line)
            i += 1
            continue
        else:
            if in_table:
                in_table = False
                headers, trows = parse_markdown_table(table_lines)
                if headers:
                    table = doc.add_table(rows=len(trows) + 1, cols=len(headers))
                    table.alignment = WD_TABLE_ALIGNMENT.CENTER
                    set_table_borders(table, color="B0C4DE", sz="6")
                    
                    # Style Header
                    hdr_cells = table.rows[0].cells
                    for col_idx, h_text in enumerate(headers):
                        cell = hdr_cells[col_idx]
                        set_cell_background(cell, "1F4E79")  # Navy Blue header
                        set_cell_margins(cell, top=140, bottom=140, left=140, right=140)
                        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                        p = cell.paragraphs[0]
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        p.paragraph_format.space_after = Pt(0)
                        clean_h = clean_latex_math(h_text.replace('**', ''))
                        run = p.add_run(clean_h)
                        run.font.name = 'Times New Roman'
                        run.font.size = Pt(9.5)
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
                    
                    # Style Data Rows
                    for r_idx, row_data in enumerate(trows):
                        row_cells = table.rows[r_idx + 1].cells
                        bg_color = "F9FAFC" if (r_idx % 2 == 1) else "FFFFFF"
                        for c_idx, cell_value in enumerate(row_data):
                            if c_idx >= len(headers):
                                continue
                            cell = row_cells[c_idx]
                            set_cell_background(cell, bg_color)
                            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
                            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                            p = cell.paragraphs[0]
                            p.paragraph_format.space_after = Pt(0)
                            
                            val_clean = clean_latex_math(cell_value)
                            if re.match(r'^-?\d+(\.\d+)?%?$', val_clean.strip()) or val_clean.strip() in ['TP', 'FP', 'TN', 'FN']:
                                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            else:
                                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                            
                            add_formatted_runs(p, cell_value)
                            for r in p.runs:
                                r.font.name = 'Times New Roman'
                                r.font.size = Pt(9.0)
                    
                    sp = doc.add_paragraph()
                    sp.paragraph_format.space_after = Pt(6)
                    sp.paragraph_format.line_spacing = 1.0
                table_lines = []
        
        # Blank line
        if not stripped:
            i += 1
            continue
            
        # Horizontal Rule
        if stripped == '---':
            i += 1
            continue
            
        # Document Title (# Title)
        if stripped.startswith('# '):
            title_text = stripped[2:].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(10)
            p.paragraph_format.line_spacing = 1.2
            run = p.add_run(clean_latex_math(title_text))
            run.font.name = 'Times New Roman'
            run.font.size = Pt(18)
            run.bold = True
            run.font.color.rgb = RGBColor(16, 44, 87)
            i += 1
            continue
            
        # Author Line
        if stripped.startswith('**A. Sai Charan, Sri Rithvik, and Krishna Kiriti**'):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run("A. Sai Charan, Sri Rithvik, and Krishna Kiriti")
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12)
            run.bold = True
            i += 1
            continue
            
        # Affiliation Lines
        if stripped.startswith('*Department of Computer Science') or stripped.startswith('*Biological Databases') or stripped.startswith('*Repository:'):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2)
            clean_affil = stripped.strip('*').strip()
            run = p.add_run(clean_affil)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(10)
            run.italic = True
            run.font.color.rgb = RGBColor(80, 80, 80)
            i += 1
            continue

        # Heading 2 (## Section Title)
        if stripped.startswith('## '):
            h_text = stripped[3:].strip()
            
            if h_text.lower() == 'abstract':
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(14)
                p.paragraph_format.space_after = Pt(4)
                run = p.add_run("ABSTRACT")
                run.font.name = 'Times New Roman'
                run.font.size = Pt(11)
                run.bold = True
                run.font.color.rgb = RGBColor(16, 44, 87)
            else:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_before = Pt(16)
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.keep_with_next = True
                run = p.add_run(clean_latex_math(h_text))
                run.font.name = 'Times New Roman'
                run.font.size = Pt(12.5)
                run.bold = True
                run.font.color.rgb = RGBColor(16, 44, 87)
            i += 1
            continue
            
        # Heading 3 (### Sub-section Title)
        if stripped.startswith('### '):
            h_text = stripped[4:].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(clean_latex_math(h_text))
            run.font.name = 'Times New Roman'
            run.font.size = Pt(11)
            run.bold = True
            run.font.color.rgb = RGBColor(31, 78, 121)
            i += 1
            continue
            
        # Heading 4 (#### Sub-sub-section Title)
        if stripped.startswith('#### '):
            h_text = stripped[5:].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(clean_latex_math(h_text))
            run.font.name = 'Times New Roman'
            run.font.size = Pt(10.5)
            run.bold = True
            run.italic = True
            run.font.color.rgb = RGBColor(60, 60, 60)
            i += 1
            continue
            
        # Table captions (TABLE I, TABLE II, etc.)
        if stripped.startswith('TABLE ') or stripped.startswith('TABLE I') or stripped.startswith('TABLE II'):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(clean_latex_math(stripped))
            run.font.name = 'Times New Roman'
            run.font.size = Pt(10)
            run.bold = True
            run.font.color.rgb = RGBColor(16, 44, 87)
            i += 1
            continue

        # Figure captions (*Fig. 1...*)
        if stripped.startswith('*Fig. ') or stripped.startswith('_Fig. '):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(10)
            clean_cap = stripped.strip('*_').strip()
            run = p.add_run(clean_latex_math(clean_cap))
            run.font.name = 'Times New Roman'
            run.font.size = Pt(9.5)
            run.italic = True
            run.font.color.rgb = RGBColor(60, 60, 60)
            i += 1
            continue
            
        # Block math equations ($$...$$)
        if stripped.startswith('$$') and stripped.endswith('$$') and len(stripped) > 4:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            eq_text = clean_latex_math(stripped[2:-2].strip())
            run = p.add_run(eq_text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(10.5)
            run.italic = True
            run.bold = False
            run.font.color.rgb = RGBColor(10, 10, 50)
            i += 1
            continue
            
        # Bullet list items (* item or - item)
        if (stripped.startswith('* ') or stripped.startswith('- ')) and not stripped.startswith('*Fig.'):
            bullet_text = stripped[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            
            if bullet_text.startswith('[x]') or bullet_text.startswith('[X]'):
                c_run = p.add_run("☑  ")
                c_run.font.name = 'Arial'
                c_run.font.size = Pt(10)
                c_run.bold = True
                c_run.font.color.rgb = RGBColor(24, 120, 24)
                bullet_text = bullet_text[3:].strip()
            elif bullet_text.startswith('[ ]'):
                c_run = p.add_run("☐  ")
                c_run.font.name = 'Arial'
                c_run.font.size = Pt(10)
                bullet_text = bullet_text[3:].strip()
                
            add_formatted_runs(p, bullet_text)
            i += 1
            continue

        # Numbered list items (1. item, 2. item, etc.)
        num_match = re.match(r'^(\d+)\.\s+(.*)$', stripped)
        if num_match:
            item_num = num_match.group(1)
            item_text = num_match.group(2)
            
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.12
            
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.first_line_indent = Inches(-0.3)
            
            n_run = p.add_run(f"[{item_num}]  ")
            n_run.font.name = 'Times New Roman'
            n_run.font.size = Pt(10)
            n_run.bold = True
            add_formatted_runs(p, item_text)
            i += 1
            continue

        # Standard Body Paragraph
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        
        # Abstract & Index Terms
        if stripped.startswith('Biomedical knowledge discovery relies') or 'Biomedical knowledge discovery' in stripped:
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.right_indent = Inches(0.4)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run("Abstract—")
            run.bold = True
            run.italic = False
            add_formatted_runs(p, stripped, italic_all=True)
        elif stripped.startswith('**Index Terms**') or 'Index Terms—' in stripped:
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.right_indent = Inches(0.4)
            p.paragraph_format.space_after = Pt(12)
            add_formatted_runs(p, stripped)
        else:
            add_formatted_runs(p, stripped)
            
        i += 1

    # Save to docs and root
    out_dir = 'docs'
    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, 'Annotation_Contradiction_Atlas_Research_Paper.docx')
    doc.save(out_path)
    
    # Also save in root directory for convenience
    root_path = 'Annotation_Contradiction_Atlas_Research_Paper.docx'
    doc.save(root_path)
    
    print(f"Generated docs path: {out_path} ({os.path.getsize(out_path)} bytes)")
    print(f"Generated root path: {root_path} ({os.path.getsize(root_path)} bytes)")
    return out_path

if __name__ == '__main__':
    create_document()
